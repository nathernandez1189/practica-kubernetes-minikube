from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
TITLE = "Práctica Kubernetes con Minikube"
SUBTITLE = "Deployments, Services, Ingress y rolling updates en ARM64"
STUDENT = "Natalia Hernández"
DATE = "9 de septiembre de 2026"

SECTIONS = [
    (
        "1. Objetivo",
        [
            "Comprender Kubernetes mediante un clúster local de un nodo, despliegues, servicios, observación, escalado, balanceo, Ingress y actualizaciones progresivas.",
        ],
    ),
    (
        "2. Entorno y adaptaciones ARM64",
        [
            "La práctica se ejecutó en servidorUbuntu, una máquina Ubuntu 22.04 ARM64 con dos CPU y cerca de 4 GB de memoria. Docker 29.8.0 ya estaba instalado.",
            "Se instalaron Minikube 1.39.0 y kubectl 1.37.0 para ARM64. Minikube inició Kubernetes 1.37.0 con containerd 2.3.4 y el controlador Docker. La aplicación Node usa node:22-alpine en lugar de Node 4.4, una versión obsoleta.",
        ],
    ),
    (
        "3. Clúster y hello-minikube",
        [
            "El nodo minikube alcanzó estado Ready con IP 192.168.49.2. Se registraron la información del clúster, los pods del sistema, eventos y la configuración activa de kubectl.",
            "hello-minikube se desplegó con kicbase/echo-server:1.0 y un Service NodePort en el puerto 30081. La respuesta identificó el pod que atendió la solicitud y se verificaron sus registros.",
        ],
    ),
    (
        "4. Dashboard y acceso",
        [
            "Se habilitó el complemento Kubernetes Dashboard. Un proxy temporal en 192.168.100.3:8001 respondió HTTP 200. El proxy se cerró al terminar.",
            "hello-node se expuso con NodePort 30080. También se ejecutó kubectl port-forward sobre el Service en 192.168.100.3:8090 y la aplicación devolvió correctamente el nombre del pod y la versión v1.",
        ],
    ),
    (
        "5. Aplicación hello-node",
        [
            "La imagen se construyó dentro del runtime de Minikube. Kubernetes desplegó hello-node:v1 y se comprobaron pods, Deployment, Service, descripción, registros, variables de entorno, código fuente y una petición ejecutada dentro del contenedor.",
        ],
    ),
    (
        "6. Escalado y balanceo",
        [
            "El Deployment pasó de una a cuatro réplicas, cada una con una IP diferente. Doce peticiones sucesivas fueron atendidas por los cuatro nombres de pod, lo que demostró el balanceo del Service.",
            "Después se redujo el Deployment a dos réplicas. Kubernetes mantuvo dos pods Running y terminó los otros dos.",
        ],
    ),
    (
        "7. Ingress Controller",
        [
            "Un Ingress Controller observa recursos Ingress y configura un proxy de entrada para dirigir tráfico HTTP o HTTPS hacia Services internos.",
            "Se habilitó Ingress NGINX y se creó el host practica.local. La ruta /node respondió desde hello-node y /echo respondió desde hello-minikube usando la misma IP de entrada.",
        ],
    ),
    (
        "8. Rolling update",
        [
            "Como desafío se construyeron hello-node:v1 y hello-node:v2. Con cuatro réplicas activas, kubectl set image inició una actualización progresiva. Kubernetes sustituyó los pods por etapas y mantuvo instancias disponibles.",
            "El historial registró dos revisiones. Ocho peticiones posteriores respondieron con v2 desde diferentes pods, confirmando la actualización y el balanceo.",
        ],
    ),
    (
        "9. Material de GPU",
        [
            "El ZIP adicional usa NVIDIA CUDA, NumPy y Numba. Una prueba real requiere una GPU NVIDIA visible, sus controladores Linux y NVIDIA Container Toolkit. VirtualBox sobre este Mac Apple Silicon no expone ese dispositivo, por lo que los archivos se conservaron como referencia y no se alteró Docker con una instalación que no podría usar aceleración.",
        ],
    ),
    (
        "10. Limpieza y restauración",
        [
            "Se eliminaron los Deployments, Services e Ingress creados. Solo quedó el Service interno kubernetes. El proxy del Dashboard se cerró y Minikube quedó completamente detenido, conservando su perfil.",
            "Los cuatro contenedores que estaban activos antes de la práctica se reiniciaron. Flask + MySQL y Flask + Redis respondieron con estado ok, por lo que el trabajo anterior quedó restaurado.",
        ],
    ),
    (
        "11. Conclusión",
        [
            "La práctica demostró el ciclo completo de una aplicación Kubernetes: construcción de imagen, despliegue, exposición, inspección, escalado, balanceo, entrada HTTP, actualización gradual y limpieza segura.",
        ],
    ),
]

RESULTS = [
    ("Nodo", "Ready", "Kubernetes 1.37.0 ARM64"),
    ("hello-minikube", "HTTP correcto", "NodePort 30081"),
    ("Dashboard", "HTTP 200", "Proxy temporal 8001"),
    ("Port-forward", "Respuesta v1", "Puerto 8090"),
    ("Escalado", "4 pods", "4 IP diferentes"),
    ("Balanceo", "4 pods observados", "12 peticiones"),
    ("Ingress", "2 rutas correctas", "/node y /echo"),
    ("Rolling update", "v1 → v2", "4 réplicas disponibles"),
    ("Limpieza", "Minikube detenido", "Contenedores restaurados"),
]


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    props.append(node)


def cell_text(cell, value, bold=False, white=False):
    cell.text = ""
    run = cell.paragraphs[0].add_run(value)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(9.2)
    run.font.color.rgb = RGBColor(255, 255, 255) if white else RGBColor(35, 48, 65)


def docx_table(doc, headers, rows, header_color="123B5D"):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    for i, value in enumerate(headers):
        shade(table.rows[0].cells[i], header_color)
        cell_text(table.rows[0].cells[i], value, bold=True, white=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell_text(cells[i], str(value))
    return table


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor(35, 48, 65)
    normal.paragraph_format.space_after = Pt(7)
    for name, size, rgb in [("Title", 29, (18, 59, 93)), ("Heading 1", 16, (18, 59, 93))]:
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(*rgb)

    header = section.header.paragraphs[0]
    header.text = "KUBERNETES · MINIKUBE · ARM64"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(37, 99, 235)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(TITLE)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(SUBTITLE)
    run.bold = True
    run.font.size = Pt(14.5)
    run.font.color.rgb = RGBColor(37, 99, 235)
    doc.add_paragraph("")
    info = docx_table(doc, ["Dato", "Valor"], [
        ("Estudiante", STUDENT),
        ("Fecha", DATE),
        ("Plataforma", "Ubuntu 22.04 · ARM64 · Docker"),
        ("Versiones", "Minikube 1.39.0 · Kubernetes/kubectl 1.37.0"),
    ])
    info.alignment = 1
    doc.add_paragraph("")
    status = doc.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    result = status.add_run("PRÁCTICA COMPLETA · INGRESS Y ROLLING UPDATE VERIFICADOS")
    result.bold = True
    result.font.color.rgb = RGBColor(5, 150, 105)
    doc.add_page_break()

    doc.add_heading("Arquitectura", level=1)
    docx_table(doc, ["Capa", "Elemento", "Función"], [
        ("Host", "servidorUbuntu", "VM ARM64 192.168.100.3"),
        ("Clúster", "Minikube", "Control plane de un nodo"),
        ("Entrada", "Ingress NGINX", "Rutas /node y /echo"),
        ("Aplicaciones", "hello-node / echo-server", "Deployments y Services"),
    ])
    doc.add_heading("Resultados verificados", level=1)
    docx_table(doc, ["Prueba", "Resultado", "Detalle"], RESULTS, "1D4ED8")

    for heading, paragraphs in SECTIONS:
        doc.add_heading(heading, level=1)
        for text in paragraphs:
            doc.add_paragraph(text)

    doc.add_heading("Referencias oficiales", level=1)
    for url in [
        "https://minikube.sigs.k8s.io/docs/start/",
        "https://minikube.sigs.k8s.io/docs/drivers/docker/",
        "https://kubernetes.io/docs/tasks/tools/install-kubectl-linux/",
        "https://kubernetes.io/docs/tasks/access-application-cluster/ingress-minikube/",
        "https://kubernetes.io/docs/tutorials/kubernetes-basics/update/update-intro/",
        "https://docs.nvidia.com/datacenter/cloud-native/container-toolkit/latest/install-guide.html",
    ]:
        doc.add_paragraph(url, style="List Bullet")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Natalia Hernández · Práctica Kubernetes con Minikube · 2026")
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(100, 116, 139)
    path = DOCS / "Informe-Practica-Kubernetes-Minikube.docx"
    doc.save(path)
    return path


def pdf_table(headers, rows, widths, cell_style, head_style, color="#123B5D"):
    data = [[Paragraph(x, head_style) for x in headers]]
    data += [[Paragraph(str(x), cell_style) for x in row] for row in rows]
    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(color)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#AFC3D1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EFF6FF")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    return table


def build_pdf():
    path = DOCS / "Informe-Practica-Kubernetes-Minikube.pdf"
    base = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "Cover", parent=base["Title"], fontName="Helvetica-Bold", fontSize=27,
        leading=33, textColor=colors.HexColor("#123B5D"), alignment=TA_CENTER,
    )
    subtitle_style = ParagraphStyle(
        "Subtitle", parent=base["Heading2"], fontSize=14, leading=19,
        textColor=colors.HexColor("#2563EB"), alignment=TA_CENTER,
    )
    heading = ParagraphStyle(
        "Heading", parent=base["Heading1"], fontName="Helvetica-Bold", fontSize=15,
        leading=19, spaceBefore=10, spaceAfter=6, textColor=colors.HexColor("#123B5D"),
    )
    body = ParagraphStyle(
        "Body", parent=base["BodyText"], fontName="Helvetica", fontSize=9.2,
        leading=13.2, spaceAfter=6.5, textColor=colors.HexColor("#233041"),
    )
    cell = ParagraphStyle("Cell", parent=body, fontSize=8.1, leading=9.7)
    head = ParagraphStyle("Head", parent=cell, fontName="Helvetica-Bold", textColor=colors.white)

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawString(2 * cm, 1.2 * cm, "Natalia Hernández · Kubernetes con Minikube")
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Página {document.page}")
        canvas.restoreState()

    pdf = SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=1.7 * cm, bottomMargin=1.8 * cm, title=TITLE, author=STUDENT,
    )
    story = [
        Spacer(1, 3.1 * cm), Paragraph(TITLE, title_style), Spacer(1, 0.35 * cm),
        Paragraph(SUBTITLE, subtitle_style), Spacer(1, 1.1 * cm),
        pdf_table(["Dato", "Valor"], [
            ("Estudiante", STUDENT),
            ("Fecha", DATE),
            ("Plataforma", "Ubuntu 22.04 · ARM64 · Docker"),
            ("Versiones", "Minikube 1.39.0 · Kubernetes/kubectl 1.37.0"),
        ], [4 * cm, 10 * cm], cell, head),
        Spacer(1, 1 * cm),
        Paragraph("PRÁCTICA COMPLETA · INGRESS Y ROLLING UPDATE VERIFICADOS", subtitle_style),
        PageBreak(),
        Paragraph("Arquitectura", heading),
        pdf_table(["Capa", "Elemento", "Función"], [
            ("Host", "servidorUbuntu", "VM ARM64 192.168.100.3"),
            ("Clúster", "Minikube", "Control plane de un nodo"),
            ("Entrada", "Ingress NGINX", "Rutas /node y /echo"),
            ("Aplicaciones", "hello-node / echo-server", "Deployments y Services"),
        ], [3.5 * cm, 5 * cm, 6 * cm], cell, head),
        Paragraph("Resultados verificados", heading),
        pdf_table(["Prueba", "Resultado", "Detalle"], RESULTS,
                  [4.1 * cm, 4.8 * cm, 5.6 * cm], cell, head, "#1D4ED8"),
    ]
    for section_title, paragraphs in SECTIONS:
        story.append(Paragraph(section_title, heading))
        for text in paragraphs:
            story.append(Paragraph(text, body))
    story.append(Paragraph("Referencias oficiales", heading))
    for url in [
        "https://minikube.sigs.k8s.io/docs/start/",
        "https://minikube.sigs.k8s.io/docs/drivers/docker/",
        "https://kubernetes.io/docs/tasks/tools/install-kubectl-linux/",
        "https://kubernetes.io/docs/tasks/access-application-cluster/ingress-minikube/",
        "https://kubernetes.io/docs/tutorials/kubernetes-basics/update/update-intro/",
        "https://docs.nvidia.com/datacenter/cloud-native/container-toolkit/latest/install-guide.html",
    ]:
        story.append(Paragraph(f"• {url}", body))
    pdf.build(story, onFirstPage=footer, onLaterPages=footer)
    return path


if __name__ == "__main__":
    print(build_docx())
    print(build_pdf())

